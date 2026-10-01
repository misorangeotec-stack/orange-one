"""
build_ink_expiry_status.py — every ink lot IN STOCK, and whether Tally has its expiry date.

Output is meant to go to the accountant: the 'To Update' sheet lists in-stock ink lots with no
expiry in Tally, with a blank column to write the date into.

    Stock      = fetch_stock_ageing_lotwise.build (Tally Operating System/tools): lot balances netted
                 from rpt_batch_line, tied to Tally's closing qty per item (oldest lots trimmed).
    Expiry     = build_ink_expiry_list (this folder): batch_expiry_raw on any voucher line of the lot.
    Ink        = group path contains 'INK' (not CHEMICALS), same rule as build_ink_expiry_list.

Stock that no lot accounts for ('(no lot record)') cannot carry an expiry — it is listed on its own
sheet. READ-ONLY — only SELECTs from ConnectWave.

    python tools/build_ink_expiry_status.py
    python tools/build_ink_expiry_status.py --out some\\path.xlsx --today 2026-09-28
"""
from __future__ import annotations

import argparse
import sys
import tempfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, r"C:/Users/admin/Desktop/Tally Operating System/tools")
import fetch_stock_ageing_lotwise as ageing  # noqa: E402
import build_ink_expiry_list as expiry  # noqa: E402

DEFAULT_OUT = Path(r"C:/Users/admin/OneDrive - Orange O Tec Private Limited/Claude Code Working File"
                   r"/Ink Lot Expiry Status.xlsx")
NO_LOT = "(no lot record)"


def company_names() -> dict[str, str]:
    """Full Tally company name (what the ageing tool reports) → short name, via the live tenant."""
    out = {}
    for c in expiry.get("v_company?select=tenant_id,company_guid,company_name"):
        if "~" not in c["tenant_id"]:
            out[c["company_name"]] = expiry.COMPANY.get(c["company_guid"], c["company_name"])
    return out


def ink_items() -> set[tuple[str, str]]:
    """(short company, item) that ANY book files under an ink group. EPN SUBLIMATION MAGENTA is
    '(Ungrouped)' in FY 2026-27 and only the FY 2024-26 book still calls it ink — same rule as the
    Hub page (frontend/src/apps/ink-expiry/lib/expiry.ts)."""
    rows = expiry.get_all("rpt_stock_summary_item", {
        "group_path": "ilike.*INK*", "select": "tenant_id,company_guid,item,group_path"}, "tenant_id,item,fy")
    co = lambda r: (expiry.COMPANY.get(r["company_guid"], r["company_guid"]), r["item"])
    out = {co(r) for r in rows if expiry.is_ink(r["group_path"])}
    # Excluded if ANY book files the item under provision / dead / diff / loose stock.
    return out - {co(r) for r in rows if expiry.is_excluded(r["group_path"], r["item"])}


def build(today: date):
    short = company_names()
    ink = ink_items()
    with tempfile.TemporaryDirectory() as tmp:
        stock, _ = ageing.build(today.strftime("%Y%m%d"), str(Path(tmp) / "ageing.xlsx"))
    lines, items, bal = expiry.fetch()
    exp_rows, _ = expiry.build(lines, items, bal, today)
    exp = {(r["company"], r["item"], r["lot"]): r for r in exp_rows}

    lots, no_lot = [], []
    for s in stock:
        co = short.get(s["Company"], s["Company"])
        if (co, s["Item Name"]) not in ink or s["Total Qty"] <= 0.001:
            continue
        base = {
            "company": co, "item": s["Item Name"], "category": s["Stock Group"],
            "main_group": s["Primary Group"], "path": s["Group Path"], "lot": s["Lot / Batch No"],
            "qty": s["Total Qty"], "uom": s["UOM"], "godown": s["Godown"],
            "inward": datetime.strptime(s["Oldest Purchase"], "%d-%b-%y").date() if s["Oldest Purchase"] else None,
            "age": s["Age (Days)"] if s["Age (Days)"] != "" else None,
        }
        if s["Lot / Batch No"] == NO_LOT:
            no_lot.append(base)
            continue
        e = exp.get((co, s["Item Name"], s["Lot / Batch No"]))
        base.update({
            "mfd": e["mfd"] if e else None,
            "exp": e["exp"] if e else None,
            "exp_raw": e["exp_raw"] if e else None,
            "days": e["days"] if e else None,
            "state": ("Not updated" if not e
                      else "Updated — not a valid date" if not e["exp"]
                      else "Updated — EXPIRED" if e["days"] < 0
                      else "Updated"),
        })
        lots.append(base)
    key = lambda r: (r["company"], r["category"] or "", r["item"], r["inward"] or date.min, r["lot"])
    return sorted(lots, key=key), sorted(no_lot, key=key)


def write_excel(lots, no_lot, out: Path, today: date):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(bold=True, color="FFFFFF")
    fill_for = {
        "Not updated": PatternFill("solid", fgColor="FCE4D6"),
        "Updated — EXPIRED": PatternFill("solid", fgColor="F8CBAD"),
        "Updated — not a valid date": PatternFill("solid", fgColor="D9D9D9"),
    }
    entry = PatternFill("solid", fgColor="FFF2CC")
    DF = "DD-MMM-YYYY"

    def sheet(ws, headers, data, widths, date_cols=(), qty_col=None):
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in data:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for col in date_cols:
            for c in ws[get_column_letter(col)][1:]:
                c.number_format = DF
        if qty_col:
            for c in ws[get_column_letter(qty_col)][1:]:
                c.number_format = "#,##0.###"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 32

    # 1. For the accountant: in-stock lots with no expiry (plus the one with a non-date expiry)
    todo = [r for r in lots if r["state"] in ("Not updated", "Updated — not a valid date")]
    ws = wb.active
    ws.title = "To Update"
    sheet(ws, ["Sr", "Company", "Category (stock group)", "Stock item", "Lot no.", "Qty in stock", "Unit",
               "Godown", "First inward date", "Age (days)", "What is in Tally now",
               "Expiry date to enter", "Updated by / on"],
          [[i, r["company"], r["category"], r["item"], r["lot"], r["qty"], r["uom"], r["godown"], r["inward"],
            r["age"], r["exp_raw"] or "(blank)", None, None] for i, r in enumerate(todo, 1)],
          [5, 17, 30, 40, 22, 11, 6, 16, 12, 8, 16, 16, 18], date_cols=(9, 12), qty_col=6)
    for row in ws.iter_rows(min_row=2, min_col=12, max_col=13):
        for c in row:
            c.fill = entry

    # 2. Everything: every in-stock ink lot, updated or not
    ws = wb.create_sheet("All In-Stock Lots")
    sheet(ws, ["Company", "Category (stock group)", "Main group", "Stock item", "Lot no.", "Qty in stock", "Unit",
               "Godown", "First inward date", "Age (days)", "Expiry status", "Mfg date", "Expiry date",
               "Days to expiry", "Group path"],
          [[r["company"], r["category"], r["main_group"], r["item"], r["lot"], r["qty"], r["uom"], r["godown"],
            r["inward"], r["age"], r["state"], r["mfd"], r["exp"] or r["exp_raw"], r["days"], r["path"]]
           for r in lots],
          [17, 30, 16, 40, 22, 11, 6, 16, 12, 8, 22, 12, 12, 9, 50], date_cols=(9, 12, 13), qty_col=6)
    for i, r in enumerate(lots, 2):
        f = fill_for.get(r["state"])
        if f:
            for c in ws[i]:
                c.fill = f

    # 3. Stock with no lot at all — needs a lot before it can take an expiry
    ws = wb.create_sheet("Stock Without Lot")
    sheet(ws, ["Company", "Category (stock group)", "Stock item", "Qty in stock", "Unit", "Godown"],
          [[r["company"], r["category"], r["item"], r["qty"], r["uom"], r["godown"]] for r in no_lot],
          [17, 30, 40, 11, 6, 16], qty_col=4)

    # 4. Summary
    ws = wb.create_sheet("Summary")
    agg: dict[tuple, Counter] = defaultdict(Counter)
    qty: dict[tuple, Counter] = defaultdict(Counter)
    for r in lots:
        k = (r["company"], r["category"])
        agg[k][r["state"]] += 1
        qty[k][r["state"]] += r["qty"]
    states = ["Updated", "Updated — EXPIRED", "Not updated", "Updated — not a valid date"]
    sheet(ws, ["Company", "Category (stock group)", "Lots in stock", *(f"{s} (lots)" for s in states),
               "Qty not updated", "% lots updated"],
          [[c, g, sum(a.values()), *(a[s] for s in states), round(qty[(c, g)]["Not updated"], 2),
            round(100 * (a["Updated"] + a["Updated — EXPIRED"]) / sum(a.values()), 0)]
           for (c, g), a in sorted(agg.items(), key=lambda kv: (kv[0][0], kv[0][1] or ""))],
          [17, 34, 9, 10, 12, 11, 12, 12, 9])
    ws.auto_filter.ref = None
    n = Counter(r["state"] for r in lots)
    r0 = ws.max_row + 2
    notes = [
        f"As at {today:%d-%b-%Y}. Source: ConnectWave (Tally sync). Ink only — stock groups whose path contains 'INK'.",
        f"{len(lots)} ink lots in stock: {n['Updated'] + n['Updated — EXPIRED']} have an expiry in Tally, "
        f"{n['Not updated']} do not ({n['Updated — EXPIRED']} of the updated ones are already past expiry).",
        "Stock per lot is tied to Tally's closing qty per item; stock no lot accounts for is on 'Stock Without Lot'.",
        "'To Update' is the sheet for the accountant — fill 'Expiry date to enter', then update the lot in Tally.",
    ]
    for i, t in enumerate(notes):
        ws.cell(r0 + i, 1, t)

    wb.move_sheet("Summary", offset=-3)
    wb.active = 0
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(out)
    except PermissionError:
        out = out.with_name(f"{out.stem}_{datetime.now():%H%M%S}{out.suffix}")
        wb.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--today", type=lambda s: datetime.strptime(s, "%Y-%m-%d").date(), default=date.today())
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")

    lots, no_lot = build(a.today)
    out = write_excel(lots, no_lot, a.out, a.today)
    print("in-stock ink lots:", len(lots), dict(Counter(r["state"] for r in lots)))
    print("by company:", {c: dict(Counter(r["state"] for r in lots if r["company"] == c))
                          for c in sorted({r["company"] for r in lots})})
    print("stock without lot:", len(no_lot), "rows,", round(sum(r["qty"] for r in no_lot)), "qty")
    print("wrote", out)


if __name__ == "__main__":
    main()

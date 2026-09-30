"""
build_ink_stabilisation_schedule.py — ink lots manufactured at Enterprises Surat, with their retest dates.

QC retests every manufactured ink lot three times: 3, 6 and 9 months after its production date.
This lists every lot of finished ink made on a STOCK JOURNAL-PRODUCTION voucher at Enterprises Surat
(both books: FY 2025-26 pre-split and FY 2026-27) and works out the three test dates.

    Production date  = the production voucher's date (first one, if a lot was made on two vouchers).
                       Not the lot name: '26061113' reads as June 2026 but its voucher is 10-Jul-2026.
    Mfg / Expiry     = the Tally batch's own dates, where the batch has them. Most lots do NOT:
                       Tally only started carrying expiry on Surat batches in Aug/Sep 2026. A blank
                       means "not entered in Tally" — it is never guessed.
    Test 1 / 2 / 3   = production date + 3 / 6 / 9 calendar months (31st rolls back to month end)

READ-ONLY. It only SELECTs from ConnectWave. Credentials come from the ConnectWave .env beside
connectwave.exe; the service key is never printed.

    python tools/build_ink_stabilisation_schedule.py                       # writes the Excel
    python tools/build_ink_stabilisation_schedule.py --out some\\path.xlsx --today 2026-09-26
"""
from __future__ import annotations

import argparse
import calendar
import json
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from datetime import date, datetime
from pathlib import Path

ENV_FILE = Path(r"C:/Users/admin/OneDrive - Orange O Tec Private Limited/AI-Sync/connectwave-deploy/.env")
DEFAULT_OUT = Path(r"C:/Users/admin/OneDrive - Orange O Tec Private Limited/Claude Code Working File"
                   r"/Ink Stabilisation - Enterprises Surat.xlsx")
SURAT_GUID = "59a6c2d9-0c5a-4fc5-b8c5-3be6fec3289e"
PRODUCTION = "STOCK JOURNAL-PRODUCTION"
TEST_MONTHS = (3, 6, 9)
PAGE = 1000
# Enterprises Surat manufactures ink, so its finished ink sits under FINISHED GOODS, not PRINTING INK.
INK_PREFIXES = ("FINISHED GOODS > FINISHED INK KGS", "PRINTING INK")
NOT_INK = ("FINISHED GOODS > FINISHED INK KGS > CHEMICALS",)
NON_LOT = {"", "primary batch", "any", "not applicable", "none"}


def load_env() -> tuple[str, str]:
    env: dict[str, str] = {}
    for line in ENV_FILE.read_text(encoding="utf-8-sig").splitlines():
        if "=" in line and not line.lstrip().startswith("#"):
            k, v = line.split("=", 1)
            env[k.strip()] = v.strip().strip('"')
    return env["SUPABASE_URL"].rstrip("/"), env["SUPABASE_SERVICE_KEY"]


URL, KEY = load_env()


def get(path: str, tries: int = 4):
    req = urllib.request.Request(f"{URL}/rest/v1/{path}", headers={"apikey": KEY, "Authorization": f"Bearer {KEY}"})
    for attempt in range(tries):
        try:
            with urllib.request.urlopen(req, timeout=180) as r:
                return json.load(r)
        except (urllib.error.HTTPError, urllib.error.URLError) as e:
            if attempt == tries - 1:
                body = e.read()[:200] if isinstance(e, urllib.error.HTTPError) else b""
                raise RuntimeError(f"{e} on {path[:120]}… {body!r}") from None
            time.sleep(3 * (attempt + 1))


def get_all(table: str, params: dict, order: str) -> list:
    """PostgREST caps a page at 1000 rows; a stable ORDER is what keeps paging from skipping rows."""
    out, off = [], 0
    while True:
        q = urllib.parse.urlencode({**params, "order": order, "limit": PAGE, "offset": off}, safe=",.*:()>")
        rows = get(f"{table}?{q}")
        out.extend(rows)
        if len(rows) < PAGE:
            return out
        off += len(rows)


def ymd(s: str) -> date:
    return datetime.strptime(s, "%Y%m%d").date()


def tally_date(s: str | None) -> date | None:
    """Tally batch dates arrive as text: '21-Aug-27', occasionally '21-Aug-2027'."""
    if not s:
        return None
    for fmt in ("%d-%b-%y", "%d-%b-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s.strip(), fmt).date()
        except ValueError:
            pass
    return None


def add_months(d: date, n: int) -> date:
    m = d.month - 1 + n
    y, m = d.year + m // 12, m % 12 + 1
    return date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


UNGROUPED = "(Ungrouped)"
INK_WORDS = ("INK", "SUBLIMATION", "REACTIVE")


def is_ink(item: str, path: str | None) -> bool:
    # A few inks have no stock group in Tally (EPN SUBLIMATION MAGENTA, the '-G' sublimation set…);
    # keep those by name, which also leaves out the ungrouped CLEANER.
    if path in (None, UNGROUPED):
        return any(w in item.upper() for w in INK_WORDS)
    return path.startswith(INK_PREFIXES) and not path.startswith(NOT_INK)


def fetch():
    base = {"company_guid": f"eq.{SURAT_GUID}"}
    prod = get_all("rpt_batch_line", {
        **base, "voucher_type": f"eq.{PRODUCTION}", "movement": "eq.in", "affects_stock": "is.true",
        "select": "tenant_id,voucher_guid,line_no,batch_no,vch_date,voucher_no,stock_item,batch_name,qty,uom,godown_name",
    }, "tenant_id,voucher_guid,line_no,batch_no")
    dated = get_all("rpt_batch_line", {
        **base, "or": "(batch_expiry_raw.not.is.null,batch_mfd.not.is.null)",
        "select": "tenant_id,voucher_guid,line_no,batch_no,vch_date,stock_item,batch_name,batch_mfd,batch_expiry_raw",
    }, "tenant_id,voucher_guid,line_no,batch_no")
    groups = get_all("rpt_stock_summary_item", {**base, "select": "tenant_id,item,group_path"}, "tenant_id,item")
    return prod, dated, groups


def build(prod, dated, groups, today: date):
    # Group path per item; the live (un-suffixed) tenant wins over older-year snapshots.
    path: dict[str, str] = {}
    for g in sorted(groups, key=lambda g: "~" not in g["tenant_id"]):
        path[g["item"]] = g["group_path"]

    # Batch dates as Tally holds them; the latest-dated line wins if a batch was edited over time.
    bdates: dict[tuple, tuple] = {}
    for r in sorted(dated, key=lambda r: r["vch_date"]):
        k = (r["stock_item"], r["batch_name"])
        mfd, exp = bdates.get(k, (None, None))
        bdates[k] = (tally_date(r["batch_mfd"]) or mfd, tally_date(r["batch_expiry_raw"]) or exp)

    lots: dict[tuple, dict] = {}
    skipped = Counter()
    for r in prod:
        grp = path.get(r["stock_item"])
        if not is_ink(r["stock_item"], grp):
            skipped[f"{grp or '(no group)'} | {r['stock_item']}" if grp in (None, UNGROUPED) else grp] += 1
            continue
        grp = grp or UNGROUPED
        k = (r["stock_item"], (r["batch_name"] or "").strip())
        d = ymd(r["vch_date"])
        lot = lots.setdefault(k, {"item": k[0], "lot": k[1], "group": grp, "prod": d, "qty": 0.0,
                                  "uom": r["uom"], "vouchers": set(), "godowns": set()})
        lot["prod"] = min(lot["prod"], d)
        lot["qty"] += r["qty"] or 0
        lot["vouchers"].add(r["voucher_no"])
        if r["godown_name"]:
            lot["godowns"].add(r["godown_name"])

    out = []
    for (item, lot_no), L in lots.items():
        mfd, exp = bdates.get((item, lot_no), (None, None))
        tests = [add_months(L["prod"], m) for m in TEST_MONTHS]
        nxt = next(((i + 1, t) for i, t in enumerate(tests) if t >= today), None)
        done_by_date = sum(t < today for t in tests)
        stage = ("Cycle complete (all 3 dates passed)" if nxt is None
                 else f"Awaiting Test {nxt[0]}" + (f" — {done_by_date} earlier date(s) passed" if done_by_date else ""))
        out.append({
            **L, "mfd": mfd, "exp": exp, "tests": tests, "stage": stage,
            "next_no": nxt[0] if nxt else None, "next_date": nxt[1] if nxt else None,
            "days": (nxt[1] - today).days if nxt else None,
            "real_lot": lot_no.lower() not in NON_LOT,
            "family": grp_leaf(L["group"]),
        })
    out.sort(key=lambda x: (x["next_date"] or date.max, x["prod"], x["item"], x["lot"]))
    return out, skipped


def grp_leaf(p: str) -> str:
    return "(No stock group in Tally)" if p == UNGROUPED else p.split(" > ")[-1].strip()


def write_excel(lots, out: Path, today: date):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(bold=True, color="FFFFFF")
    soon = PatternFill("solid", fgColor="FCE4D6")   # next test within 30 days
    month = PatternFill("solid", fgColor="FFF2CC")  # next test within 31-60 days
    grey = Font(color="808080")
    DF = "DD-MMM-YYYY"

    def sheet(ws, headers, rows, widths, date_cols=()):
        ws.append(headers)
        for c in ws[1]:
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for col in date_cols:
            for c in ws[get_column_letter(col)][1:]:
                c.number_format = DF
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 32

    # 1. One row per lot
    ws = wb.active
    ws.title = "Retest Schedule"
    rows = [[L["item"], L["family"], L["lot"], L["prod"], L["mfd"], L["exp"], round(L["qty"], 2), L["uom"],
             *L["tests"], L["next_no"] and f"Test {L['next_no']}", L["next_date"], L["days"], L["stage"],
             ", ".join(sorted(L["vouchers"])), ", ".join(sorted(L["godowns"]))] for L in lots]
    sheet(ws, ["Stock item", "Ink family", "Lot no.", "Production date", "Mfg date (Tally)", "Expiry date (Tally)",
               "Qty produced", "Unit", "Test 1 (+3 m)", "Test 2 (+6 m)", "Test 3 (+9 m)", "Next test",
               "Next test date", "Days to next test", "Stage", "Production voucher(s)", "Godown"],
          rows, [38, 26, 16, 13, 13, 13, 11, 6, 13, 13, 13, 9, 13, 10, 34, 18, 18], date_cols=(4, 5, 6, 9, 10, 11, 13))
    for r in range(2, ws.max_row + 1):
        days = ws.cell(r, 14).value
        if days is not None and days <= 30:
            for c in ws[r]:
                c.fill = soon
        elif days is not None and days <= 60:
            for c in ws[r]:
                c.fill = month
        elif days is None:
            for c in ws[r]:
                c.font = grey

    # 2. One row per test — the QC calendar, from this month on
    ws = wb.create_sheet("Test Calendar")
    start = today.replace(day=1)
    cal = sorted(((t, n + 1, L) for L in lots for n, t in enumerate(L["tests"]) if t >= start),
                 key=lambda x: (x[0], x[2]["item"], x[2]["lot"]))
    sheet(ws, ["Test date", "Month", "Test", "Stock item", "Ink family", "Lot no.", "Production date",
               "Expiry date (Tally)", "Qty produced", "Done on", "Result / remarks"],
          [[t, t.strftime("%b %Y"), f"Test {n}", L["item"], L["family"], L["lot"], L["prod"], L["exp"],
            round(L["qty"], 2), None, None] for t, n, L in cal],
          [13, 10, 8, 38, 26, 16, 13, 13, 11, 13, 30], date_cols=(1, 7, 8, 10))

    # 3. Summary: tests per month for the next 12 months, by test number
    ws = wb.create_sheet("Summary")
    months = [add_months(start, i) for i in range(12)]
    count = Counter((t.replace(day=1), n) for t, n, _ in cal)
    sheet(ws, ["Month", "Test 1", "Test 2", "Test 3", "Total lots to test"],
          [[m.strftime("%b %Y"), count[(m, 1)], count[(m, 2)], count[(m, 3)],
            sum(count[(m, i)] for i in (1, 2, 3))] for m in months], [12, 9, 9, 9, 16])
    ws.auto_filter.ref = None
    r = ws.max_row + 2
    with_exp = sum(1 for L in lots if L["exp"])
    notes = [
        f"Source: ConnectWave (Tally sync) — Enterprises Surat, {PRODUCTION} vouchers, inward finished-ink lines.",
        f"Built {datetime.now():%d-%b-%Y %H:%M} with 'today' = {today:%d-%b-%Y}.",
        f"{len(lots)} lots; production dates {min(L['prod'] for L in lots):%d-%b-%Y} to {max(L['prod'] for L in lots):%d-%b-%Y}.",
        f"Expiry date is entered in Tally for only {with_exp} of {len(lots)} lots. Blank = not in Tally (not guessed).",
        "Tests fall at production date + 3, + 6 and + 9 months. Dates in the past are shown as passed — ",
        "ConnectWave has no record of whether a test was actually done; use 'Done on' in Test Calendar.",
        "Row colours on Retest Schedule: orange = next test within 30 days; yellow = within 60; grey = all 3 passed.",
    ]
    for i, n in enumerate(notes):
        ws.cell(r + i, 1, n)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--today", type=lambda s: datetime.strptime(s, "%Y-%m-%d").date(), default=date.today())
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8")

    prod, dated, groups = fetch()
    lots, skipped = build(prod, dated, groups, a.today)
    write_excel(lots, a.out, a.today)

    print(f"production inward lines: {len(prod)}; ink lots: {len(lots)}")
    print("skipped (not ink):", dict(skipped.most_common()))
    print("lots without a real lot no.:", sum(not L["real_lot"] for L in lots))
    print("lots with Tally expiry:", sum(1 for L in lots if L["exp"]))
    print("stage:", dict(Counter(L["stage"].split(" —")[0] for L in lots)))
    print("wrote", a.out)


if __name__ == "__main__":
    main()
